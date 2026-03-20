# Small web app: resume matcher + basic user accounts + uploads
from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    session,
    send_from_directory,
    send_file,
)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from flask_sqlalchemy import SQLAlchemy
import re
import os
import datetime
from io import BytesIO
import docx

BASE_DIR = os.path.dirname(__file__)
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///app.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
app.secret_key = "change-me-to-a-random-secret"

db = SQLAlchemy(app)

# Allowed upload extensions
ALLOWED_EXTENSIONS = {"pdf", "png", "jpg", "jpeg", "doc", "docx", "txt"}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


# Predefined skill keywords (core database for analyzer)
SKILLS = [
    # Programming Languages
    "python","java","c","c++","c#","javascript","typescript","go","rust","ruby","php","kotlin","swift","scala","r","matlab","dart","objective-c",
    # Web Development
    "html","css","sass","less","bootstrap","tailwind css","react","reactjs","angular","vue","nodejs","express","nextjs","nuxtjs","jquery","web development","frontend development","backend development","full stack development",
    # Mobile Development
    "android development","ios development","flutter","react native","xamarin","ionic","mobile app development","kotlin android","swift ios",
    # Frameworks
    "django","flask","spring","spring boot","laravel","symfony","rails","asp.net","fastapi","nestjs","svelte","ember",
    # Databases
    "sql","mysql","postgresql","mongodb","sqlite","oracle","redis","cassandra","firebase","dynamodb","mariadb","neo4j",
    # Data Science & AI
    "data science","data analysis","data analytics","machine learning","deep learning","artificial intelligence","nlp","natural language processing","computer vision","reinforcement learning","predictive modeling",
    # Data Tools
    "pandas","numpy","scikit-learn","tensorflow","pytorch","keras","matplotlib","seaborn","plotly","statsmodels",
    # Big Data
    "hadoop","spark","pyspark","hive","pig","kafka","flink","big data analytics","data warehousing","etl","data pipelines",
    # BI & Visualization
    "excel","advanced excel","power bi","tableau","looker","qlikview","data visualization","dashboard development",
    # Cloud Computing
    "aws","amazon web services","azure","google cloud","gcp","cloud computing","serverless","cloud architecture",
    # DevOps
    "docker","kubernetes","terraform","ansible","jenkins","ci/cd","github actions","gitlab ci","devops","infrastructure as code",
    # Version Control
    "git","github","gitlab","bitbucket","version control",
    # Operating Systems
    "linux","unix","windows server","bash","shell scripting","command line",
    # APIs
    "rest api","graphql","soap","api development","api integration","microservices","web services",
    # Testing
    "unit testing","integration testing","automation testing","selenium","cypress","jest","pytest","software testing","qa testing","manual testing",
    # Security
    "cybersecurity","ethical hacking","penetration testing","network security","application security","cryptography","vulnerability assessment","security auditing",
    # Networking
    "networking","tcp/ip","dns","http","https","vpn","firewalls","network troubleshooting",
    # UI/UX
    "ui design","ux design","figma","adobe xd","wireframing","prototyping","user research","interaction design",
    # Design Tools
    "photoshop","illustrator","canva","indesign","graphic design","visual design",
    # Project Management
    "project management","agile","scrum","kanban","jira","confluence","waterfall methodology","risk management","stakeholder management",
    # Product & Business
    "product management","business analysis","requirements gathering","market research","product strategy","roadmap planning",
    # Finance & Accounting
    "financial analysis","budgeting","forecasting","accounting","taxation","auditing","cost management",
    # Marketing
    "digital marketing","seo","search engine optimization","sem","google analytics","social media marketing","content marketing","email marketing","affiliate marketing","ppc advertising",
    # Sales
    "sales strategy","lead generation","crm","customer acquisition","negotiation","business development","sales forecasting",
    # Customer Support
    "customer service","technical support","help desk","client relations","customer success",
    # HR Skills
    "recruitment","talent acquisition","employee relations","payroll management","performance management","hr analytics",
    # Education & Training
    "teaching","training","curriculum development","instructional design","e-learning","mentoring",
    # Research
    "research methods","data interpretation","statistical analysis","academic research",
    # Soft Skills
    "communication","verbal communication","written communication","teamwork","collaboration","leadership","problem solving","critical thinking","analytical thinking","decision making","adaptability","creativity","innovation","time management","organization","attention to detail","presentation skills","public speaking","negotiation","conflict resolution","emotional intelligence","interpersonal skills","active listening","multitasking","strategic thinking","planning","self motivation","work ethic","stress management",
    # Management Skills
    "people management","team leadership","performance coaching","delegation","resource management","change management",
    # Office Tools
    "microsoft office","word","excel","powerpoint","outlook","google workspace","google sheets","google docs","slack","trello","notion",
    # Emerging Technologies
    "blockchain","cryptocurrency","web3","iot","internet of things","robotics","augmented reality","virtual reality","edge computing","quantum computing",
    # Miscellaneous Technical
    "debugging","troubleshooting","system design","software architecture","design patterns","object oriented programming","data structures","algorithms","optimization","code review","technical documentation"
]

# resume builder templates (display names)
TEMPLATES = [
    "Modern Professional Template",
    "Clean Minimal Template",
    "Executive Format",
    "Corporate ATS Layout",
    "Tech Developer Format",
    "Fresher Resume Format",
    "Two-Column Modern Template",
    "Functional Style Template",
    "Skills-Based Template",
    "Creative Designer Template",
    "Simple Chronological Template",
    "Graduate Student Template",
    "Project-Focused Template",
    "Software Engineer Template",
    "Data Analyst Template",
    "Managerial Layout",
    "Academic CV Style",
    "Internship Resume Format",
    "One-Page ATS Template",
    "Classic Professional Layout",
]

# mapping from template name to actual text content with placeholders
TEMPLATE_TEXTS = {
    "Modern Professional Template":
        "Name: {name}\nEducation: {education}\nSkills: {skills}\nExperience: {experience}\nProjects: {projects}\nCertifications: {certifications}\n",
    "Clean Minimal Template":
        "{name}\n{education}\n{skills}\n{experience}\n{projects}\n{certifications}\n",
    "Executive Format":
        "EXECUTIVE PROFILE\nName: {name}\nEducation: {education}\nKey Skills: {skills}\nExperience: {experience}\n",
    "Corporate ATS Layout":
        "{name} | {education} | {skills}\nExperience:\n{experience}\nProjects:\n{projects}\nCertifications:\n{certifications}\n",
    # ... other templates can use similar placeholder patterns
    "Tech Developer Format":
        "Developer Resume\n{name}\nEducation: {education}\nSkills:\n{skills}\nExperience:\n{experience}\n",
    # for brevity, use same base for remaining templates
}
# make sure every template has an entry (generic fallback)
for name in TEMPLATES:
    TEMPLATE_TEXTS.setdefault(
        name,
        f"{name}\nName: {{name}}\nEducation: {{education}}\nSkills: {{skills}}\nExperience: {{experience}}\nProjects: {{projects}}\nCertifications: {{certifications}}\n"
    )

# demo resume data for template previews and the editor prefill
DEMO_RESUME_DATA = {
    'name': 'Diya Agarwal',
    'email': 'd.agarwal@example.in',
    'phone': '+91 11 5555 3345',
    'summary': 'Customer-focused Retail Sales professional with solid understanding of retail dynamics, 5 years of experience delivering strong customer satisfaction and revenue growth.',
    'education': 'Diploma in Financial Accounting | Oxford Software Institute',
    'skills': 'Cash register operation, Sales expertise, Inventory management, Documentation, Teamwork',
    'experience': 'Retail Sales Associate - ZARA | New Delhi (2017 - Present)\nBarista - Dunkin\' Donuts | New Delhi (2015 - 2017)',
    'projects': 'Improved point-of-sale speed by 18%, created customer engagement plan, optimized inventory data.',
    'certifications': 'Diploma in Financial Accounting',
    'languages': 'Hindi (Native), English (Fluent), Bengali (Intermediate)'
}


SUSPICIOUS_PHRASES = [
    "registration fee", "registration fees", "guaranteed job", "earn money fast",
    "work from home and earn", "pay to apply", "send money", "western union",
    "wire transfer", "no interview", "get rich", "training fee", "cash reward",
]

SUSPICIOUS_PHRASES = [
    "registration fee", "registration fees", "guaranteed job", "earn money fast",
    "work from home and earn", "pay to apply", "send money", "western union",
    "wire transfer", "no interview", "get rich", "training fee", "cash reward",
]


def extract_skills(text: str):
    """Extract skills from text by checking predefined keywords."""
    if not text:
        return set()
    t = text.lower()
    found = set()
    for skill in SKILLS:
        if skill in t:
            found.add(skill)
    # also pick out comma-separated tokens in explicit lists
    tokens = re.split(r"[\n,;:\t()]+", text.lower())
    for token in tokens:
        token = token.strip()
        if len(token) > 1 and token in SKILLS:
            found.add(token)
    return found


def detect_fake_phrases(text: str):
    if not text:
        return []
    t = text.lower()
    flags = []
    for phrase in SUSPICIOUS_PHRASES:
        if phrase in t:
            flags.append(phrase)
    return flags


def estimate_years_of_experience(text: str):
    if not text:
        return 0
    matches = re.findall(r"(\d+)\s*\+?\s*years?", text.lower())
    values = [int(v) for v in matches if v.isdigit()]
    if values:
        return max(values)
    return 0


def extract_education_level(text: str):
    if not text:
        return ""
    level_map = [
        (r"phd", 5), (r"doctor", 5), (r"master", 4), (r"mba", 4), (r"m\.tech", 4), (r"ms", 4),
        (r"bachelor", 3), (r"b\.tech", 3), (r"ba", 3), (r"bs", 3), (r"degree", 2),
        (r"diploma", 2), (r"high school", 1)
    ]
    text_lower = text.lower()
    for keyword, level in level_map:
        if re.search(r"\b" + keyword + r"\b", text_lower):
            return level
    return 0


def compute_resume_strength(resume_text: str, resume_skills: set):
    if not resume_text:
        return 0
    points = 0
    total = 5
    # presence heuristics
    if resume_skills:
        points += 1
    if re.search(r"\bexperience\b|\d+\s+years?", resume_text, re.I):
        points += 1
    if re.search(r"\b(bachelor|master|mba|phd|diploma|college|university)\b", resume_text, re.I):
        points += 1
    if re.search(r"\b(project|initiative|deliverable)\b", resume_text, re.I):
        points += 1
    if re.search(r"\b(certification|certified|training)\b", resume_text, re.I):
        points += 1
    return int((points / total) * 100)


def extract_keywords(text: str):
    if not text:
        return set()
    words = re.findall(r"[a-zA-Z0-9]+", text.lower())
    stopwords = set(["and","or","the","is","a","an","to","in","on","for","with","of","by","at","from","as","that","this","it","be","are","was","were"])
    return set(w for w in words if w not in stopwords and len(w) > 2)


def extract_project_relevance(text: str):
    if not text:
        return 0
    project_terms = re.findall(r"\b(project|projects|project-based|project experience|developed)\b", text.lower())
    return min(5, len(project_terms))


def compute_project_score(resume_count: int, job_count: int):
    if job_count > 0:
        return int(min(100, (resume_count / job_count) * 100))
    return 50 if resume_count > 0 else 0


def compute_ats_compatibility(resume_text: str):
    if not resume_text:
        return 0
    score = 0
    if re.search(r"\b(contact|email|phone|address|linkedin)\b", resume_text, re.I):
        score += 20
    if re.search(r"\b(skills|technical skills|skills:)\b", resume_text, re.I):
        score += 20
    if re.search(r"\b(experience|work history|professional experience)\b", resume_text, re.I):
        score += 20
    if re.search(r"\b(education|degree|bachelor|master|diploma|phd)\b", resume_text, re.I):
        score += 20
    if re.search(r"\b(project|achievement|certification|training)\b", resume_text, re.I):
        score += 20
    return min(100, score)


def resume_match_level(score: int):
    if score >= 80:
        return "Strong Match"
    if score >= 55:
        return "Moderate Match"
    return "Low Match"


def extract_text_from_file(path):
    # support docx and plain text
    if not path:
        return ""
    ext = path.rsplit('.', 1)[-1].lower()
    try:
        if ext == 'docx':
            doc = docx.Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        else:
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
    except Exception:
        return ""


# simple slug helper for templates

def slugify(text: str):
    return re.sub(r'[^a-z0-9]+', '-', text.lower()).strip('-')


# PDF generator for submitted resume data
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

def generate_pdf(data: dict, template_text: str = None) -> bytes:
    """Creates a basic PDF given the resume fields.

    If ``template_text`` is provided it is assumed to contain placeholders
    like `{name}`, `{education}` etc.  Those are filled with values from
    ``data`` and then the resulting text is written directly to the PDF.
    Otherwise the previous default layout is used.
    """
    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=letter)
    width, height = letter
    x_margin = 72
    y = height - 72

    if template_text:
        filled = template_text.format(
            name=data.get('name', ''),
            education=data.get('education', ''),
            skills=data.get('skills', ''),
            experience=data.get('experience', ''),
            projects=data.get('projects', ''),
            certifications=data.get('certifications', ''),
        )
        p.setFont("Helvetica", 11)
        for line in filled.split('\n'):
            p.drawString(x_margin, y, line)
            y -= 14
            if y < 72:
                p.showPage()
                y = height - 72
                p.setFont("Helvetica", 11)
    else:
        # header
        p.setFont("Helvetica-Bold", 18)
        p.drawString(x_margin, y, data.get('name', ''))
        y -= 28

        p.setFont("Helvetica", 12)
        # sections
        for label in ['education', 'skills', 'experience', 'projects', 'certifications']:
            text = data.get(label)
            if text:
                p.setFont("Helvetica-Bold", 12)
                p.drawString(x_margin, y, label.capitalize() + ":")
                y -= 16
                p.setFont("Helvetica", 11)
                for line in str(text).split('\n'):
                    p.drawString(x_margin + 12, y, line)
                    y -= 14
                    if y < 72:  # new page
                        p.showPage()
                        y = height - 72
                        p.setFont("Helvetica", 11)
                y -= 10

    p.showPage()
    p.save()
    pdf = buffer.getvalue()
    buffer.close()
    return pdf


# --- Database models ---


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(120))
    email = db.Column(db.String(200), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)
    is_admin = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.DateTime, default=datetime.datetime.utcnow)

    resumes = db.relationship("Resume", backref="user", lazy=True)
    applications = db.relationship("Application", backref="user", lazy=True)

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)


class Resume(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    filename = db.Column(db.String(300))
    text = db.Column(db.Text)
    uploaded_at = db.Column(db.DateTime, default=datetime.datetime.utcnow)


class Application(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    company = db.Column(db.String(250))
    status = db.Column(db.String(50), default="Applied")
    created_at = db.Column(db.DateTime, default=datetime.datetime.utcnow)


with app.app_context():
    db.create_all()
    # create a default admin if none exists
    if not User.query.filter_by(is_admin=True).first():
        admin = User(name="Admin", email="admin@example.com", is_admin=True)
        admin.set_password("change-me")
        db.session.add(admin)
        db.session.commit()

# --- End DB models ---


@app.route("/", methods=["GET"]) 
def index():
    return render_template("index.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        name = request.form.get("name")
        email = request.form.get("email")
        password = request.form.get("password")
        if not email or not password:
            flash("Email and password required", "error")
            return redirect(url_for("register"))
        if User.query.filter_by(email=email).first():
            flash("User already exists, please login", "error")
            return redirect(url_for("login"))
        u = User(name=name, email=email)
        u.set_password(password)
        db.session.add(u)
        db.session.commit()
        flash("Account created — please login", "success")
        return redirect(url_for("login"))
    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email")
        password = request.form.get("password")
        u = User.query.filter_by(email=email).first()
        if u and u.check_password(password):
            session["user_id"] = u.id
            session["is_admin"] = bool(u.is_admin)
            flash("Logged in", "success")
            return redirect(url_for("dashboard"))
        flash("Invalid credentials", "error")
        return redirect(url_for("login"))
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("Logged out", "info")
    return redirect(url_for("index"))


@app.route("/dashboard")
def dashboard():
    uid = session.get("user_id")
    if not uid:
        return redirect(url_for("login"))
    user = User.query.get(uid)
    resumes = user.resumes
    applications = user.applications
    return render_template("dashboard.html", user=user, resumes=resumes, applications=applications)


@app.route("/upload_resume", methods=["POST"])
def upload_resume():
    uid = session.get("user_id")
    if not uid:
        flash("Login required", "error")
        return redirect(url_for("login"))
    user = User.query.get(uid)

    # either file or pasted text
    file = request.files.get("file")
    text = request.form.get("resume_text")

    filename = None
    if file and file.filename and allowed_file(file.filename):
        filename = secure_filename(f"{user.id}_{int(datetime.datetime.utcnow().timestamp())}_" + file.filename)
        path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
        file.save(path)

    r = Resume(user_id=user.id, filename=filename, text=text)
    db.session.add(r)
    db.session.commit()
    flash("Resume uploaded", "success")
    return redirect(url_for("dashboard"))



# --- Resume builder ---

@app.route('/resume_builder', methods=['GET'])
def resume_builder():
    uid = session.get('user_id')
    if not uid:
        return redirect(url_for('login'))

    return render_template('resume_builder.html',
                           templates=TEMPLATES,
                           template_texts=TEMPLATE_TEXTS,
                           demo_data=DEMO_RESUME_DATA)


@app.route('/resume_builder/edit', methods=['GET', 'POST'])
def edit_resume():
    uid = session.get('user_id')
    if not uid:
        return redirect(url_for('login'))

    promo_template = request.args.get('template') or request.args.get('template_name') or (TEMPLATES[0] if TEMPLATES else '')
    selected_template = promo_template
    template_content = TEMPLATE_TEXTS.get(selected_template, '')
    data = {}
    rendered_template = ''
    preview = False

    if request.method == 'POST':
        skills_list = request.form.getlist('skills')
        data = {
            'name': request.form.get('name',''),
            'email': request.form.get('email',''),
            'phone': request.form.get('phone',''),
            'summary': request.form.get('summary',''),
            'education': request.form.get('education',''),
            'skills': ', '.join(skills_list) if skills_list else '',
            'skills_list': skills_list,
            'experience': request.form.get('experience',''),
            'projects': request.form.get('projects',''),
            'certifications': request.form.get('certifications',''),
            'languages': request.form.get('languages',''),
            'resume_theme': request.form.get('resume_theme','blue')
        }
        selected_template = request.form.get('template_name', selected_template)
        template_content = request.form.get('template_content', template_content)

        try:
            rendered_template = template_content.format(
                name=data.get('name',''),
                email=data.get('email',''),
                phone=data.get('phone',''),
                summary=data.get('summary',''),
                education=data.get('education',''),
                skills=data.get('skills',''),
                experience=data.get('experience',''),
                projects=data.get('projects',''),
                certifications=data.get('certifications',''),
                languages=data.get('languages','')
            )
        except Exception:
            rendered_template = template_content

        action = request.form.get('action')
        if action == 'download':
            pdf_bytes = generate_pdf(data, template_text=rendered_template)
            return send_file(BytesIO(pdf_bytes), as_attachment=True, download_name='resume.pdf', mimetype='application/pdf')

        preview = True
    else:
        # prefill with demo data for first load
        data = DEMO_RESUME_DATA.copy()
        rendered_template = template_content.format(
            name=data.get('name',''),
            email=data.get('email',''),
            phone=data.get('phone',''),
            summary=data.get('summary',''),
            education=data.get('education',''),
            skills=data.get('skills',''),
            experience=data.get('experience',''),
            projects=data.get('projects',''),
            certifications=data.get('certifications',''),
            languages=data.get('languages','')
        )
        data['skills_list'] = [s.strip() for s in data.get('skills','').split(',') if s.strip()]
        data['resume_theme'] = 'blue'

    return render_template('resume_builder_edit.html',
                           templates=TEMPLATES,
                           template_texts=TEMPLATE_TEXTS,
                           demo_data=DEMO_RESUME_DATA,
                           selected_template=selected_template,
                           template_content=template_content,
                           data=data,
                           rendered_template=rendered_template,
                           preview=preview)


@app.route("/uploads/<path:filename>")
def uploaded_file(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename)


@app.route("/admin")
def admin():
    if not session.get("is_admin"):
        flash("Admin access required", "error")
        return redirect(url_for("login"))
    users = User.query.order_by(User.created_at.desc()).all()
    resumes = Resume.query.order_by(Resume.uploaded_at.desc()).all()
    applications = Application.query.order_by(Application.created_at.desc()).all()
    return render_template("admin.html", users=users, resumes=resumes, applications=applications)


@app.route("/analyze", methods=["POST"])
def analyze():
    # support selecting an uploaded resume, file upload, or pasted resume text
    resume_text = ""
    resume_id = request.form.get("resume_id")
    if resume_id:
        r = Resume.query.get(int(resume_id))
        if r:
            # prefer file text when available
            if r.filename:
                resume_text = extract_text_from_file(os.path.join(app.config["UPLOAD_FOLDER"], r.filename))
            resume_text = (resume_text or r.text or "")

    # file upload on-the-fly
    file = request.files.get("resume_file")
    if file and file.filename and allowed_file(file.filename):
        tmp = os.path.join(app.config["UPLOAD_FOLDER"], "_tmp_" + secure_filename(file.filename))
        file.save(tmp)
        resume_text = extract_text_from_file(tmp) or resume_text
        try:
            os.remove(tmp)
        except Exception:
            pass

    # pasted resume text fallback
    resume_text = resume_text or request.form.get("resume") or request.form.get("resume_text") or ""
    job_text = request.form.get("job") or ""

    resume_skills = extract_skills(resume_text)
    job_skills = extract_skills(job_text)

    matched = resume_skills.intersection(job_skills)
    missing = job_skills.difference(resume_skills)

    skill_match_pct = int(len(matched) / len(job_skills) * 100) if len(job_skills) > 0 else 0

    resume_experience_years = estimate_years_of_experience(resume_text)
    job_experience_years = estimate_years_of_experience(job_text)
    if job_experience_years > 0:
        experience_match_pct = min(100, int((resume_experience_years / job_experience_years) * 100))
    else:
        experience_match_pct = 100 if resume_experience_years > 0 else 50

    resume_edu_level = extract_education_level(resume_text)
    job_edu_level = extract_education_level(job_text)
    if job_edu_level > 0:
        education_match_pct = int(min(100, (resume_edu_level / job_edu_level) * 100))
    else:
        education_match_pct = 100 if resume_edu_level > 0 else 50

    resume_keywords = extract_keywords(resume_text)
    job_keywords = extract_keywords(job_text)
    if len(job_keywords) > 0:
        keyword_match_pct = int(len(resume_keywords.intersection(job_keywords)) / len(job_keywords) * 100)
    else:
        keyword_match_pct = 0

    project_hits = extract_project_relevance(resume_text)
    job_project_hits = extract_project_relevance(job_text)
    project_relevance_pct = compute_project_score(project_hits, job_project_hits)

    resume_strength = compute_resume_strength(resume_text, resume_skills)
    ats_score = compute_ats_compatibility(resume_text)

    job_fit_score = int(
        skill_match_pct * 0.50 +
        experience_match_pct * 0.20 +
        education_match_pct * 0.10 +
        keyword_match_pct * 0.10 +
        project_relevance_pct * 0.10
    )

    fit_category = resume_match_level(job_fit_score)

    flags = detect_fake_phrases(job_text)
    if len(flags) == 0:
        risk = "Low"
    elif len(flags) <= 2:
        risk = "Medium"
    else:
        risk = "High"

    section_complete = {
        'summary': bool(re.search(r"\b(summary|objective)\b", resume_text, re.I)),
        'skills': bool(resume_skills),
        'experience': bool(re.search(r"\b(experience|work\s+history|professional\s+experience)\b", resume_text, re.I)),
        'education': bool(re.search(r"\b(education|degree|bachelor|master|phd|diploma)\b", resume_text, re.I)),
        'projects': bool(project_hits)
    }

    career_readiness = int((resume_strength + job_fit_score + ats_score) / 3)
    recommended_skills = sorted(job_skills.difference(resume_skills))

    return render_template(
        "result.html",
        score=job_fit_score,
        matched=sorted(matched),
        missing=sorted(missing),
        resume_skills=sorted(resume_skills),
        job_skills=sorted(job_skills),
        risk=risk,
        flags=flags,
        skill_match_pct=skill_match_pct,
        experience_match_pct=experience_match_pct,
        education_match_pct=education_match_pct,
        keyword_match_pct=keyword_match_pct,
        project_relevance_pct=project_relevance_pct,
        resume_strength=resume_strength,
        ats_score=ats_score,
        career_readiness=career_readiness,
        fit_category=fit_category,
        section_complete=section_complete,
        recommended_skills=recommended_skills,
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
